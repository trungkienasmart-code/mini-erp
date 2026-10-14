# app/modules/backup/services.py
import os
import zipfile
import shutil
import json
from datetime import datetime, timezone, timedelta
from pathlib import Path
from typing import List, Optional, Tuple

from sqlmodel import Session, select, func

from app.core.config import BASE_DIR


# ============================================================
# CONSTANTS
# ============================================================
BACKUP_DIR = BASE_DIR / "backups"
UPLOAD_DIR = BASE_DIR / "app" / "static" / "uploads"
DB_FILE = BASE_DIR / "database.db"
SECRET_FILE = BASE_DIR / ".secret_key"

BACKUP_DIR.mkdir(parents=True, exist_ok=True)


# ============================================================
# TẠO BACKUP
# ============================================================
def create_backup(note: str = "", user_id: Optional[int] = None) -> Tuple[Optional[Path], str]:
    """
    Tạo file backup zip chứa:
    - database.db
    - File uploads (logo, internal)
    - .secret_key
    - metadata.json
    """
    try:
        if not DB_FILE.exists():
            return None, "File database.db không tồn tại"

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_name = f"backup_{timestamp}.zip"
        backup_path = BACKUP_DIR / backup_name

        with zipfile.ZipFile(backup_path, 'w', zipfile.ZIP_DEFLATED) as zipf:
            # 1. Database
            zipf.write(DB_FILE, arcname="database.db")

            # 2. File uploads
            if UPLOAD_DIR.exists():
                for file_path in UPLOAD_DIR.rglob("*"):
                    if file_path.is_file():
                        rel = file_path.relative_to(BASE_DIR)
                        zipf.write(file_path, arcname=str(rel))

            # 3. Secret key
            if SECRET_FILE.exists():
                zipf.write(SECRET_FILE, arcname=".secret_key")

            # 4. Metadata
            meta = {
                "created_at": datetime.now(timezone.utc).isoformat(),
                "note": note,
                "user_id": user_id,
                "version": "1.0.0",
                "db_size": DB_FILE.stat().st_size,
            }
            zipf.writestr("metadata.json", json.dumps(meta, ensure_ascii=False, indent=2))

        size_mb = backup_path.stat().st_size / 1024 / 1024
        return backup_path, f"Đã tạo backup ({size_mb:.1f} MB)"

    except Exception as e:
        return None, f"Lỗi tạo backup: {e}"


def list_backups() -> List[dict]:
    """Liệt kê tất cả file backup, mới nhất trước."""
    if not BACKUP_DIR.exists():
        return []

    backups = []
    for f in sorted(BACKUP_DIR.glob("backup_*.zip"), reverse=True):
        stat = f.stat()

        # Đọc metadata (nếu có)
        note = ""
        try:
            with zipfile.ZipFile(f, 'r') as zipf:
                if "metadata.json" in zipf.namelist():
                    meta = json.loads(zipf.read("metadata.json"))
                    note = meta.get("note", "")
        except Exception:
            pass

        backups.append({
            "name": f.name,
            "path": str(f),
            "size": stat.st_size,
            "size_mb": stat.st_size / 1024 / 1024,
            "created_at": datetime.fromtimestamp(stat.st_mtime, tz=timezone.utc),
            "note": note,
        })

    return backups


def delete_backup(backup_name: str) -> Tuple[bool, str]:
    """Xóa file backup."""
    if ".." in backup_name or "/" in backup_name or "\\" in backup_name:
        return False, "Tên file không hợp lệ"

    path = BACKUP_DIR / backup_name
    if not path.exists() or not path.is_file():
        return False, "File backup không tồn tại"

    try:
        path.unlink()
        return True, f"Đã xóa {backup_name}"
    except Exception as e:
        return False, f"Lỗi xóa: {e}"


def restore_backup(backup_name: str) -> Tuple[bool, str]:
    """
    Phục hồi từ file backup.
    ⚠️ CẢNH BÁO: Ghi đè database + uploads hiện tại.
    """
    if ".." in backup_name or "/" in backup_name or "\\" in backup_name:
        return False, "Tên file không hợp lệ"

    backup_path = BACKUP_DIR / backup_name
    if not backup_path.exists():
        return False, "File backup không tồn tại"

    try:
        # Tạo backup an toàn của DB hiện tại trước khi ghi đè
        if DB_FILE.exists():
            safe_backup = BACKUP_DIR / f"before_restore_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db"
            shutil.copy2(DB_FILE, safe_backup)

        # Extract
        with zipfile.ZipFile(backup_path, 'r') as zipf:
            # Giải nén database
            if "database.db" in zipf.namelist():
                # Ghi đè database
                with zipf.open("database.db") as src:
                    with open(DB_FILE, 'wb') as dst:
                        shutil.copyfileobj(src, dst)

            # Giải nén uploads
            for name in zipf.namelist():
                if name.startswith("app/static/uploads/") and not name.endswith("/"):
                    target = BASE_DIR / name
                    target.parent.mkdir(parents=True, exist_ok=True)
                    with zipf.open(name) as src:
                        with open(target, 'wb') as dst:
                            shutil.copyfileobj(src, dst)

            # Giải nén secret key
            if ".secret_key" in zipf.namelist():
                with zipf.open(".secret_key") as src:
                    with open(SECRET_FILE, 'wb') as dst:
                        shutil.copyfileobj(src, dst)

        return True, "Đã phục hồi. Vui lòng ĐÓNG và MỞ LẠI ứng dụng để áp dụng."

    except Exception as e:
        return False, f"Lỗi phục hồi: {e}"


# ============================================================
# CHECK REMINDER — Nhắc backup 1 tuần/lần
# ============================================================
def get_backup_status() -> dict:
    """Trả về trạng thái backup cho nhắc nhở."""
    backups = list_backups()
    now = datetime.now(timezone.utc)

    if not backups:
        return {
            "has_backup": False,
            "days_since_last": None,
            "need_backup": True,
            "message": "Chưa có backup nào. Nên sao lưu ngay!",
            "urgency": "danger",
        }

    last = backups[0]["created_at"]
    if last.tzinfo is None:
        last = last.replace(tzinfo=timezone.utc)
    days_since = (now - last).days

    if days_since >= 7:
        return {
            "has_backup": True,
            "days_since_last": days_since,
            "need_backup": True,
            "message": f"Đã {days_since} ngày từ lần sao lưu cuối. Nên sao lưu ngay!",
            "urgency": "danger",
        }
    elif days_since >= 5:
        return {
            "has_backup": True,
            "days_since_last": days_since,
            "need_backup": False,
            "message": f"Lần sao lưu cuối cách đây {days_since} ngày.",
            "urgency": "warning",
        }
    else:
        return {
            "has_backup": True,
            "days_since_last": days_since,
            "need_backup": False,
            "message": f"Backup gần nhất: {days_since} ngày trước. Tốt!",
            "urgency": "success",
        }


# ============================================================
# EXPORT EXCEL — Từng module
# ============================================================
def export_products_excel(session: Session) -> Tuple[Optional[bytes], str]:
    """Export danh sách sản phẩm ra Excel."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment

        from app.modules.quan_ly_kho.models import Product, Category

        wb = Workbook()
        ws = wb.active
        ws.title = "Sản phẩm"

        headers = ["STT", "Mã SKU", "Tên sản phẩm", "ĐVT",
                   "Danh mục", "Giá nhập", "Giá bán", "VAT (%)",
                   "Tồn kho", "Tồn tối thiểu", "Phương pháp giá vốn"]

        # Header
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0d6efd")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Data
        products = session.exec(select(Product).order_by(Product.sku)).all()
        for idx, p in enumerate(products, 1):
            cat = session.get(Category, p.category_id) if p.category_id else None
            ws.append([
                idx, p.sku, p.name, p.unit,
                cat.name if cat else "",
                p.cost_price, p.price, p.tax_rate,
                p.stock_quantity, p.min_stock, p.cost_method,
            ])

        # Độ rộng cột
        widths = [6, 15, 30, 8, 20, 12, 12, 10, 10, 12, 15]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + i)].width = w

        # Save to bytes
        from io import BytesIO
        output = BytesIO()
        wb.save(output)
        return output.getvalue(), "OK"

    except Exception as e:
        return None, f"Lỗi export: {e}"


def export_customers_excel(session: Session) -> Tuple[Optional[bytes], str]:
    """Export danh sách khách hàng."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment

        from app.modules.ban_hang.models import Customer

        wb = Workbook()
        ws = wb.active
        ws.title = "Khách hàng"

        headers = ["STT", "Mã KH", "Tên khách hàng", "SĐT",
                   "Email", "Địa chỉ", "MST", "Loại",
                   "Hạn mức nợ", "Công nợ hiện tại", "Điểm tích lũy"]

        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="198754")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        customers = session.exec(select(Customer).order_by(Customer.code)).all()
        for idx, c in enumerate(customers, 1):
            ws.append([
                idx, c.code, c.name, c.phone,
                c.email, c.address, c.tax_code, c.customer_type,
                c.credit_limit, c.current_debt, c.loyalty_points,
            ])

        widths = [6, 12, 25, 15, 25, 30, 15, 10, 15, 15, 12]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + i)].width = w

        from io import BytesIO
        output = BytesIO()
        wb.save(output)
        return output.getvalue(), "OK"

    except Exception as e:
        return None, f"Lỗi export: {e}"


def export_suppliers_excel(session: Session) -> Tuple[Optional[bytes], str]:
    """Export danh sách nhà cung cấp."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment

        from app.modules.quan_ly_kho.models import Supplier

        wb = Workbook()
        ws = wb.active
        ws.title = "Nhà cung cấp"

        headers = ["STT", "Mã NCC", "Tên NCC", "SĐT",
                   "Email", "Địa chỉ", "MST", "Số ngày nợ", "Công nợ hiện tại"]

        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="dc3545")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        suppliers = session.exec(select(Supplier).order_by(Supplier.code)).all()
        for idx, s in enumerate(suppliers, 1):
            ws.append([
                idx, s.code, s.name, s.phone,
                s.email, s.address, s.tax_code,
                s.payment_term_days, s.current_debt,
            ])

        widths = [6, 12, 25, 15, 25, 30, 15, 12, 15]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + i)].width = w

        from io import BytesIO
        output = BytesIO()
        wb.save(output)
        return output.getvalue(), "OK"

    except Exception as e:
        return None, f"Lỗi export: {e}"


def export_orders_excel(
    session: Session,
    date_from=None, date_to=None,
) -> Tuple[Optional[bytes], str]:
    """Export đơn bán hàng theo khoảng thời gian."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from app.modules.ban_hang.models import SaleOrder, SaleOrderLine, Customer

        wb = Workbook()
        ws = wb.active
        ws.title = "Đơn bán hàng"

        headers = ["STT", "Mã đơn", "Ngày", "Khách hàng",
                   "Số SP", "Tổng tiền", "Đã trả", "Còn lại", "Trạng thái"]

        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0d6efd")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        stmt = select(SaleOrder)
        if date_from:
            stmt = stmt.where(SaleOrder.created_at >= datetime.combine(
                date_from, datetime.min.time(), tzinfo=timezone.utc))
        if date_to:
            stmt = stmt.where(SaleOrder.created_at <= datetime.combine(
                date_to, datetime.max.time(), tzinfo=timezone.utc))
        stmt = stmt.order_by(SaleOrder.created_at.desc())

        orders = session.exec(stmt).all()
        for idx, o in enumerate(orders, 1):
            cust = session.get(Customer, o.customer_id) if o.customer_id else None
            line_count = session.exec(
                select(func.count()).select_from(SaleOrderLine)
                .where(SaleOrderLine.order_id == o.id)
            ).one() or 0

            ws.append([
                idx, o.code,
                o.created_at.strftime("%d/%m/%Y %H:%M"),
                cust.name if cust else "Khách lẻ",
                line_count, o.total_amount, o.paid_amount,
                o.total_amount - o.paid_amount, o.status,
            ])

        widths = [6, 18, 18, 25, 8, 15, 15, 15, 15]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + i)].width = w

        from io import BytesIO
        output = BytesIO()
        wb.save(output)
        return output.getvalue(), "OK"

    except Exception as e:
        return None, f"Lỗi export: {e}"


def get_backup_stats() -> dict:
    """Thống kê về backups."""
    backups = list_backups()
    total_size = sum(b["size"] for b in backups) / 1024 / 1024

    return {
        "count": len(backups),
        "total_size_mb": total_size,
        "latest": backups[0] if backups else None,
    }